"""Generate the QCert User Manual as a Word document."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.paragraph_format.space_after = Pt(6)

# ── Title ──
title = doc.add_heading("QCert — User Manual", level=0)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

intro = doc.add_paragraph(
    "QCert is a desktop application for generating professional, "
    "print-ready PDF certificates in bulk. You supply a background PDF "
    "template and an Excel spreadsheet of recipient data, design the "
    "layout visually, then export individual or combined PDFs."
)
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER

doc.add_page_break()

# ── 1  Basic Workflow ──
doc.add_heading("1. Basic Workflow", level=1)

# 1.1
doc.add_heading("1.1  Load Your Data", level=2)
doc.add_paragraph(
    "Go to File > Open Spreadsheet and select an .xlsx file. "
    "The first row must contain column headers (e.g. Name, Course, Date). "
    "Every subsequent row represents one certificate recipient."
)
doc.add_paragraph(
    "Once loaded, use the Next / Previous buttons or press N / P on "
    "your keyboard to browse through records. You can also type a row "
    "number into the \"Go to row\" field or use the search function."
)

# 1.2
doc.add_heading("1.2  Set a Background Template", level=2)
doc.add_paragraph(
    "Go to Template > Set Background PDF and choose your certificate "
    "design file. Then select the page size (A4 or Letter) and "
    "orientation (Landscape or Portrait) to match your template."
)

# 1.3
doc.add_heading("1.3  Add Text Fields", level=2)
doc.add_paragraph(
    "Open the Text Fields tab and click Add. Configure each field:"
)
bullets = [
    ("Source Column", "Link to a spreadsheet column (e.g. Name, Course)."),
    ("Static Text", "Type fixed text instead (e.g. \"Certificate of Completion\")."),
    ("Combined Columns", "Merge multiple columns with a separator "
     "(e.g. First Name + Last Name)."),
    ("Position", "Set X and Y coordinates (in mm) and width, "
     "or drag the field directly on the canvas."),
    ("Style", "Choose font, size, bold/normal, colour (hex code), "
     "and alignment (left, centre, or right)."),
    ("Format Rule", "Apply automatic formatting such as uppercase, "
     "titlecase, or a date format (see Section 4)."),
]
for label, desc in bullets:
    p = doc.add_paragraph(style="List Bullet")
    run_b = p.add_run(f"{label} — ")
    run_b.bold = True
    p.add_run(desc)

# 1.4
doc.add_heading("1.4  Add Images", level=2)
doc.add_paragraph(
    "Open the Images tab and click Add. Select a PNG or JPG file "
    "(e.g. a logo, signature, or stamp). Set the position, size, "
    "rotation, and opacity."
)
doc.add_paragraph(
    "To use a different image for each certificate, link the image to a "
    "spreadsheet column that contains file paths."
)

# 1.5
doc.add_heading("1.5  Preview and Position Elements", level=2)
doc.add_paragraph(
    "The canvas updates live as you make changes. Click any element to "
    "select it, then drag it to reposition. Use the arrow keys to nudge "
    "an element by 0.5 mm, or hold Shift + Arrow for a fine nudge of 0.1 mm."
)

# 1.6
doc.add_heading("1.6  Export Certificates", level=2)
doc.add_paragraph("Open the Output tab and configure your export:")
bullets_export = [
    ("Filename template",
     "Use {ColumnName} placeholders, e.g. {Name}_{ID}."),
    ("Output mode",
     "Choose Individual PDFs (one file per record) or "
     "Combined (one multi-page PDF)."),
]
for label, desc in bullets_export:
    p = doc.add_paragraph(style="List Bullet")
    run_b = p.add_run(f"{label} — ")
    run_b.bold = True
    p.add_run(desc)

doc.add_paragraph(
    "Then click Export Current (single record), Export All, or "
    "Export Selected (specify rows, e.g. 1-5, 8, 10). "
    "A progress bar will show the status of batch exports."
)

# ── 2  Saving & Reusing Layouts ──
doc.add_heading("2. Saving and Reusing Layouts", level=1)
doc.add_paragraph(
    "Go to File > Save Layout to save your entire design — all text "
    "fields, images, positions, and output settings — as a .json file. "
    "Reload it later with File > Open Layout. QCert automatically keeps "
    "the five most recent backups of each layout."
)
doc.add_paragraph(
    "You can also duplicate an existing layout to create a variant "
    "without starting from scratch, or reset to a blank default layout."
)

# ── 3  Keyboard Shortcuts ──
doc.add_heading("3. Keyboard Shortcuts", level=1)

shortcut_data = [
    ("N", "Next record"),
    ("P", "Previous record"),
    ("Arrow keys", "Nudge selected element by 0.5 mm"),
    ("Shift + Arrow keys", "Fine nudge by 0.1 mm"),
    ("Ctrl + S", "Save layout"),
    ("Ctrl + O", "Open spreadsheet"),
    ("Ctrl + Z", "Undo"),
    ("Ctrl + Y", "Redo"),
]

table = doc.add_table(rows=1, cols=2, style="Light Shading Accent 1")
table.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr = table.rows[0].cells
hdr[0].text = "Key"
hdr[1].text = "Action"
for cell in hdr:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for key, action in shortcut_data:
    row = table.add_row().cells
    row[0].text = key
    row[1].text = action

# ── 4  Format Rules Reference ──
doc.add_heading("4. Format Rules Reference", level=1)
doc.add_paragraph(
    "When adding a text field, you can set a Format Rule to "
    "automatically transform the cell value before it is placed "
    "on the certificate."
)

format_data = [
    ("uppercase", "Converts text to ALL CAPS"),
    ("lowercase", "Converts text to all lower case"),
    ("titlecase", "Capitalises Each Word"),
    ("sentencecase", "Capitalises the first word only"),
    ("DD MMM YYYY", "Formats a date as 15 Jan 2026"),
    ("DD/MM/YYYY", "Formats a date as 15/01/2026"),
    ("YYYY-MM-DD", "Formats a date as 2026-01-15"),
    ("today", "Inserts today's date"),
    ("today:FORMAT", "Inserts today's date with a custom format (strftime)"),
    ("number:N", "Formats a number with N decimal places"),
]

table2 = doc.add_table(rows=1, cols=2, style="Light Shading Accent 1")
table2.alignment = WD_TABLE_ALIGNMENT.CENTER
hdr2 = table2.rows[0].cells
hdr2[0].text = "Rule"
hdr2[1].text = "Effect"
for cell in hdr2:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.bold = True

for rule, effect in format_data:
    row = table2.add_row().cells
    row[0].text = rule
    row[1].text = effect

# ── 5  Available Fonts ──
doc.add_heading("5. Available Fonts", level=1)

doc.add_paragraph().add_run("Built-in:").bold = True
doc.add_paragraph(
    "Helvetica, Times-Roman, Courier (each with bold and italic variants).",
    style="List Bullet",
)

doc.add_paragraph().add_run("Bundled:").bold = True
doc.add_paragraph(
    "Carlito, Inter, Open Sans, Roboto (regular, bold, and italic).",
    style="List Bullet",
)

doc.add_paragraph().add_run("System fonts:").bold = True
doc.add_paragraph(
    "QCert automatically detects fonts installed on your operating system.",
    style="List Bullet",
)

# ── 6  Building a Windows Installer (Optional) ──
doc.add_heading("6. Building a Windows Installer (Optional)", level=1)
doc.add_paragraph(
    "To distribute QCert as a standalone application that does not "
    "require Python, run:"
)
p_code = doc.add_paragraph()
run_code = p_code.add_run("    pip install pyinstaller\n    python build.py")
run_code.font.name = "Consolas"
run_code.font.size = Pt(10)

doc.add_paragraph(
    "This produces a self-contained executable. Compile the generated "
    ".iss file with Inno Setup to create a full Windows installer "
    "(no admin privileges required on the target machine)."
)

# ── Footer note ──
doc.add_paragraph("")
footer = doc.add_paragraph("QCert is fully offline — no internet connection required.")
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
for run in footer.runs:
    run.italic = True

# ── Save ──
output_path = "/home/user/QCert/QCert_User_Manual.docx"
doc.save(output_path)
print(f"Manual saved to {output_path}")
